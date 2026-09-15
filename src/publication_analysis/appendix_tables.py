"""Appendix publication tables derived from already-canonical evidence.

This module is a display/export layer only. It reads persisted publication
analysis artifacts and representation metadata; it never reruns, refits,
tunes, or recomputes any statistical model or multiple-testing result.
"""
from __future__ import annotations

import ast
from pathlib import Path
from typing import Mapping

import pandas as pd

from publication_analysis import publication_exports as exports
from publication_analysis.compact_tables import (
    TABLE2_COMPACT_COLUMNS,
    CompactTableError,
    _clean_cell,
    _format_ci,
    _stage_map_from_contract,
)

APPENDIX_TABLE_A2_COLUMNS = TABLE2_COMPACT_COLUMNS
APPENDIX_TABLE_A2_EXPECTED_SOURCE_ROWS = 81
APPENDIX_TABLE_A2_P_VALUE_NUMBER_FORMAT = "0.000000000000E+00"


class AppendixTableError(CompactTableError):
    """Raised when an appendix table cannot be derived from canonical evidence."""


def _canonical_predictor_order(representation_contract: pd.DataFrame) -> list[str]:
    if "predictor" not in representation_contract.columns:
        raise AppendixTableError("Representation contract missing 'predictor'.")
    predictors = representation_contract["predictor"].astype(str).tolist()
    if len(predictors) != APPENDIX_TABLE_A2_EXPECTED_SOURCE_ROWS:
        raise AppendixTableError(
            f"Appendix Table A2 expected "
            f"{APPENDIX_TABLE_A2_EXPECTED_SOURCE_ROWS} canonical predictors; "
            f"observed {len(predictors)}."
        )
    duplicates = sorted(
        representation_contract.loc[
            representation_contract["predictor"].duplicated(), "predictor"
        ].astype(str)
    )
    if duplicates:
        raise AppendixTableError(
            f"Representation contract contains duplicate predictors: {duplicates}"
        )
    return predictors


def _require_columns(df: pd.DataFrame, required: set[str], name: str) -> None:
    missing = sorted(required - set(df.columns))
    if missing:
        raise AppendixTableError(f"{name} missing columns: {missing}")


def _is_blank(value: object) -> bool:
    return pd.isna(value) or value == ""


def _parse_category_order(value: object, predictor: str) -> list[object]:
    if isinstance(value, list):
        return value
    if _is_blank(value):
        return []
    if isinstance(value, str):
        try:
            parsed = ast.literal_eval(value)
        except (SyntaxError, ValueError) as exc:
            raise AppendixTableError(
                f"{predictor}: category_order is not parseable: {value!r}"
            ) from exc
        if isinstance(parsed, list):
            return parsed
    raise AppendixTableError(
        f"{predictor}: category_order must be a list or list-like string."
    )


def _single_row_by_predictor(
    df: pd.DataFrame, predictor: str, *, row_type: str, name: str
) -> pd.Series:
    matches = df.loc[
        (df["predictor"].astype(str) == predictor) & (df["row_type"] == row_type)
    ]
    if len(matches) != 1:
        raise AppendixTableError(
            f"Expected one {row_type!r} row for {predictor!r} in {name}; "
            f"found {len(matches)}."
        )
    return matches.iloc[0]


def _fdr_by_predictor(hypothesis_fdr: pd.DataFrame) -> pd.DataFrame:
    _require_columns(
        hypothesis_fdr,
        {"predictor", "raw_p", "bh_q", "in_bh_family", "fit_status"},
        "Hypothesis FDR table",
    )
    if hypothesis_fdr["predictor"].duplicated().any():
        duplicates = sorted(
            hypothesis_fdr.loc[
                hypothesis_fdr["predictor"].duplicated(), "predictor"
            ].astype(str)
        )
        raise AppendixTableError(
            f"Hypothesis FDR table contains duplicate predictors: {duplicates}"
        )
    if len(hypothesis_fdr) != APPENDIX_TABLE_A2_EXPECTED_SOURCE_ROWS:
        raise AppendixTableError(
            f"Hypothesis FDR table expected "
            f"{APPENDIX_TABLE_A2_EXPECTED_SOURCE_ROWS} source-level rows; "
            f"observed {len(hypothesis_fdr)}."
        )
    return hypothesis_fdr.set_index("predictor", drop=False)


def _assert_level_detail_matches_master(
    univariable_master: pd.DataFrame,
    categorical_level_detail: pd.DataFrame,
) -> None:
    master_levels = univariable_master.loc[
        univariable_master["row_type"] == "level"
    ].reset_index(drop=True)
    detail_levels = categorical_level_detail.loc[
        categorical_level_detail["row_type"] == "level"
    ].reset_index(drop=True)
    columns = [
        "predictor",
        "row_type",
        "level",
        "reference_level",
        "level_n",
        "level_events",
        "level_non_events",
        "level_unique_subjects",
        "odds_ratio",
        "ci_low",
        "ci_high",
        "coef_p_value",
        "fit_status",
        "separation_status",
        "support_status",
        "reason",
    ]
    _require_columns(master_levels, set(columns), "Univariable master level rows")
    _require_columns(detail_levels, set(columns), "Categorical level-detail table")
    try:
        pd.testing.assert_frame_equal(
            detail_levels.loc[:, columns]
            .astype("object")
            .where(detail_levels.loc[:, columns].notna(), ""),
            master_levels.loc[:, columns]
            .astype("object")
            .where(master_levels.loc[:, columns].notna(), ""),
            check_dtype=False,
        )
    except AssertionError as exc:
        raise AppendixTableError(
            "Categorical level-detail source differs from univariable_master "
            "level rows for displayed values/status fields."
        ) from exc


def _contract_rows_by_predictor(
    representation_contract: pd.DataFrame,
) -> pd.DataFrame:
    _require_columns(
        representation_contract,
        {"predictor", "inferential_role", "reference_level", "category_order"},
        "Representation contract",
    )
    return representation_contract.set_index("predictor", drop=False)


def _expected_reportable_levels(contract_row: pd.Series) -> list[str]:
    predictor = str(contract_row["predictor"])
    reference = contract_row.get("reference_level")
    category_order = _parse_category_order(contract_row.get("category_order"), predictor)
    if not category_order:
        raise AppendixTableError(f"{predictor}: categorical predictor has empty category_order.")
    if _is_blank(reference):
        raise AppendixTableError(f"{predictor}: categorical predictor has empty reference_level.")
    reference_as_text = str(reference)
    category_order_as_text = [str(v) for v in category_order]
    if reference_as_text not in category_order_as_text:
        raise AppendixTableError(
            f"{predictor}: reference_level {reference_as_text!r} is not in "
            f"category_order {category_order_as_text!r}."
        )
    return [level for level in category_order_as_text if level != reference_as_text]


def _assert_category_levels_follow_contract(
    categorical_level_detail: pd.DataFrame,
    representation_contract: pd.DataFrame,
) -> None:
    contract_by_predictor = _contract_rows_by_predictor(representation_contract)
    categorical_predictors = contract_by_predictor.loc[
        contract_by_predictor["inferential_role"] == "CATEGORICAL", "predictor"
    ].astype(str).tolist()
    level_detail = categorical_level_detail.loc[
        categorical_level_detail["row_type"] == "level"
    ].copy()
    level_predictors = set(level_detail["predictor"].astype(str))
    unexpected = sorted(level_predictors - set(categorical_predictors))
    if unexpected:
        raise AppendixTableError(
            "Level-detail row(s) found for non-categorical/out-of-contract "
            f"predictor(s): {unexpected}"
        )

    for predictor in categorical_predictors:
        contract_row = contract_by_predictor.loc[predictor]
        expected_levels = _expected_reportable_levels(contract_row)
        observed_rows = level_detail.loc[
            level_detail["predictor"].astype(str) == predictor
        ]
        observed_levels = observed_rows["level"].astype(str).tolist()
        if observed_levels != expected_levels:
            raise AppendixTableError(
                f"{predictor}: category level order/membership mismatch. "
                f"Expected {expected_levels!r}; observed {observed_levels!r}."
            )
        observed_references = observed_rows["reference_level"].astype(str).unique().tolist()
        expected_reference = str(contract_row["reference_level"])
        if observed_references != [expected_reference]:
            raise AppendixTableError(
                f"{predictor}: reference-level mismatch. Expected every level row "
                f"to use {expected_reference!r}; observed {observed_references!r}."
            )


def _source_display_row(
    predictor: str,
    source: pd.Series,
    fdr: pd.Series,
    stage: object,
    label: str,
) -> dict[str, object]:
    return {
        "Variable": predictor,
        "Clinical Label": label,
        "Row type": "source",
        "Level": "",
        "Reference": _clean_cell(source.get("reference_level", "")),
        "Stage": _clean_cell(stage),
        "Effect representation / comparison": _clean_cell(
            source.get("effect_representation", "")
        ),
        "N": _clean_cell(source.get("model_analysis_N", "")),
        "Events": _clean_cell(source.get("model_events", "")),
        "Non-events": _clean_cell(source.get("model_non_events", "")),
        "Unique subjects": _clean_cell(source.get("model_unique_subjects", "")),
        "Crude OR": _clean_cell(source.get("odds_ratio", "")),
        "95% CI": _format_ci(source, "ci_low", "ci_high"),
        "Predictor-level raw p": _clean_cell(fdr.get("raw_p", "")),
        "Predictor-level BH q": _clean_cell(fdr.get("bh_q", "")),
        "Level p": "",
        "In BH family": _clean_cell(fdr.get("in_bh_family", "")),
        "Fit status": _clean_cell(source.get("fit_status", "")),
        "Separation status": _clean_cell(source.get("separation_status", "")),
        "Support status": _clean_cell(source.get("support_status", "")),
        "Reason": _clean_cell(source.get("reason", "")),
    }


def _level_display_row(
    predictor: str,
    level: pd.Series,
    stage: object,
    label: str,
) -> dict[str, object]:
    return {
        "Variable": predictor,
        "Clinical Label": label,
        "Row type": "category_subrow",
        "Level": _clean_cell(level.get("level", "")),
        "Reference": _clean_cell(level.get("reference_level", "")),
        "Stage": _clean_cell(stage),
        "Effect representation / comparison": _clean_cell(
            level.get("effect_representation", "")
        ),
        "N": _clean_cell(level.get("level_n", "")),
        "Events": _clean_cell(level.get("level_events", "")),
        "Non-events": _clean_cell(level.get("level_non_events", "")),
        "Unique subjects": _clean_cell(level.get("level_unique_subjects", "")),
        "Crude OR": _clean_cell(level.get("odds_ratio", "")),
        "95% CI": _format_ci(level, "ci_low", "ci_high"),
        "Predictor-level raw p": "",
        "Predictor-level BH q": "",
        "Level p": _clean_cell(level.get("coef_p_value", "")),
        "In BH family": False,
        "Fit status": _clean_cell(level.get("fit_status", "")),
        "Separation status": _clean_cell(level.get("separation_status", "")),
        "Support status": _clean_cell(level.get("support_status", "")),
        "Reason": _clean_cell(level.get("reason", "")),
    }


def build_appendix_table_a2(
    univariable_master: pd.DataFrame,
    hypothesis_fdr: pd.DataFrame,
    categorical_level_detail: pd.DataFrame,
    representation_contract: pd.DataFrame,
    predictor_display_labels: Mapping[str, str] | None = None,
) -> pd.DataFrame:
    """Return full Appendix Table A2 univariable display rows.

    Source rows are emitted once per canonical predictor, in representation
    contract order. Categorical subrows are appended underneath their source
    predictor using the already-saved categorical level-detail evidence.
    """
    predictor_display_labels = predictor_display_labels or {}
    canonical_order = _canonical_predictor_order(representation_contract)
    stage_map = _stage_map_from_contract(representation_contract)

    _require_columns(
        univariable_master,
        {"predictor", "row_type", "fit_status", "support_status"},
        "Univariable master",
    )
    _require_columns(
        categorical_level_detail,
        {"predictor", "row_type", "level", "fit_status", "support_status"},
        "Categorical level-detail table",
    )
    _assert_level_detail_matches_master(univariable_master, categorical_level_detail)
    _assert_category_levels_follow_contract(categorical_level_detail, representation_contract)
    fdr_by_predictor = _fdr_by_predictor(hypothesis_fdr)

    missing_fdr = [p for p in canonical_order if p not in fdr_by_predictor.index]
    if missing_fdr:
        raise AppendixTableError(
            f"Hypothesis FDR table missing canonical predictor(s): {missing_fdr}"
        )

    level_detail = categorical_level_detail.loc[
        categorical_level_detail["row_type"] == "level"
    ].copy()
    unexpected_level_predictors = sorted(
        set(level_detail["predictor"].astype(str)) - set(canonical_order)
    )
    if unexpected_level_predictors:
        raise AppendixTableError(
            "Categorical level-detail table contains non-canonical predictor(s): "
            f"{unexpected_level_predictors}"
        )

    rows: list[dict[str, object]] = []
    for predictor in canonical_order:
        source = _single_row_by_predictor(
            univariable_master, predictor, row_type="source", name="univariable_master"
        )
        fdr = fdr_by_predictor.loc[predictor]
        label = predictor_display_labels.get(predictor, predictor)
        rows.append(
            _source_display_row(
                predictor=predictor,
                source=source,
                fdr=fdr,
                stage=stage_map.get(predictor, ""),
                label=label,
            )
        )

        level_rows = level_detail.loc[level_detail["predictor"].astype(str) == predictor]
        for _, level in level_rows.iterrows():
            rows.append(
                _level_display_row(
                    predictor=predictor,
                    level=level,
                    stage=stage_map.get(predictor, ""),
                    label=label,
                )
            )

    appendix = pd.DataFrame(rows, columns=APPENDIX_TABLE_A2_COLUMNS)
    assert_appendix_table_a2_contract(
        appendix,
        univariable_master=univariable_master,
        hypothesis_fdr=hypothesis_fdr,
        categorical_level_detail=categorical_level_detail,
        representation_contract=representation_contract,
        predictor_display_labels=predictor_display_labels,
    )
    return appendix


def _series_value_equal(left: object, right: object) -> bool:
    if pd.isna(left) and (pd.isna(right) or right == ""):
        return True
    if left == "" and pd.isna(right):
        return True
    return left == right


def _assert_value_equal(
    observed: object, expected: object, *, row_label: str, column: str
) -> None:
    if not _series_value_equal(observed, expected):
        raise AppendixTableError(
            f"{row_label}: {column} differs from canonical source "
            f"(observed={observed!r}, expected={expected!r})."
        )


def assert_appendix_table_a2_contract(
    appendix_table_a2: pd.DataFrame,
    *,
    univariable_master: pd.DataFrame,
    hypothesis_fdr: pd.DataFrame,
    categorical_level_detail: pd.DataFrame,
    representation_contract: pd.DataFrame,
    predictor_display_labels: Mapping[str, str] | None = None,
) -> None:
    """Fail loudly if Appendix A2 violates source/order/BH invariants."""
    predictor_display_labels = predictor_display_labels or {}
    canonical_order = _canonical_predictor_order(representation_contract)
    _require_columns(
        appendix_table_a2,
        set(APPENDIX_TABLE_A2_COLUMNS),
        "Appendix Table A2",
    )
    _assert_level_detail_matches_master(univariable_master, categorical_level_detail)
    _assert_category_levels_follow_contract(categorical_level_detail, representation_contract)
    stage_map = _stage_map_from_contract(representation_contract)
    observed_source = appendix_table_a2.loc[
        appendix_table_a2["Row type"] == "source", "Variable"
    ].astype(str).tolist()
    if len(observed_source) != APPENDIX_TABLE_A2_EXPECTED_SOURCE_ROWS:
        raise AppendixTableError(
            f"Appendix Table A2 expected "
            f"{APPENDIX_TABLE_A2_EXPECTED_SOURCE_ROWS} source rows; "
            f"observed {len(observed_source)}."
        )
    if observed_source != canonical_order:
        raise AppendixTableError(
            "Appendix Table A2 source order differs from canonical order."
        )
    duplicate_sources = sorted({p for p in observed_source if observed_source.count(p) > 1})
    if duplicate_sources:
        raise AppendixTableError(
            f"Appendix Table A2 duplicate source predictor(s): {duplicate_sources}"
        )

    category_rows = appendix_table_a2.loc[
        appendix_table_a2["Row type"] == "category_subrow"
    ]
    if category_rows["Predictor-level BH q"].replace("", pd.NA).notna().any():
        raise AppendixTableError("Category subrow carries predictor-level BH q.")
    if category_rows["Predictor-level raw p"].replace("", pd.NA).notna().any():
        raise AppendixTableError("Category subrow carries predictor-level raw p.")
    if category_rows["In BH family"].astype(str).str.lower().isin(["true", "1"]).any():
        raise AppendixTableError("Category subrow is marked as part of the BH family.")

    fdr_by_predictor = _fdr_by_predictor(hypothesis_fdr)
    level_detail = categorical_level_detail.loc[
        categorical_level_detail["row_type"] == "level"
    ].copy()
    expected_total = APPENDIX_TABLE_A2_EXPECTED_SOURCE_ROWS + len(level_detail)
    if len(appendix_table_a2) != expected_total:
        raise AppendixTableError(
            f"Appendix Table A2 expected {expected_total} total rows "
            f"(81 sources + {len(level_detail)} category detail rows); "
            f"observed {len(appendix_table_a2)}."
        )

    for predictor in canonical_order:
        source_row = appendix_table_a2.loc[
            (appendix_table_a2["Row type"] == "source")
            & (appendix_table_a2["Variable"].astype(str) == predictor)
        ].iloc[0]
        source = _single_row_by_predictor(
            univariable_master, predictor, row_type="source", name="univariable_master"
        )
        fdr = fdr_by_predictor.loc[predictor]
        row_label = f"source {predictor!r}"
        expected_source_values = {
            "Clinical Label": predictor_display_labels.get(predictor, predictor),
            "Reference": source.get("reference_level", ""),
            "Stage": stage_map.get(predictor, ""),
            "Effect representation / comparison": source.get(
                "effect_representation", ""
            ),
            "N": source.get("model_analysis_N", ""),
            "Events": source.get("model_events", ""),
            "Non-events": source.get("model_non_events", ""),
            "Unique subjects": source.get("model_unique_subjects", ""),
            "Crude OR": source.get("odds_ratio", ""),
            "95% CI": _format_ci(source, "ci_low", "ci_high"),
            "Predictor-level raw p": fdr.get("raw_p", ""),
            "Predictor-level BH q": fdr.get("bh_q", ""),
            "In BH family": fdr.get("in_bh_family", ""),
            "Fit status": source.get("fit_status", ""),
            "Separation status": source.get("separation_status", ""),
            "Support status": source.get("support_status", ""),
            "Reason": source.get("reason", ""),
        }
        for column, expected in expected_source_values.items():
            _assert_value_equal(
                source_row[column], expected, row_label=row_label, column=column
            )

    for level_index, level in level_detail.iterrows():
        predictor = str(level["predictor"])
        matches = category_rows.loc[
            (category_rows["Variable"].astype(str) == predictor)
            & (category_rows["Level"].astype(str) == str(level.get("level", "")))
            & (category_rows["Reference"].astype(str) == str(level.get("reference_level", "")))
        ]
        if len(matches) != 1:
            raise AppendixTableError(
                f"Level-detail row {level_index} for {predictor!r} is missing or duplicated "
                f"in Appendix Table A2; found {len(matches)} matches."
            )
        display = matches.iloc[0]
        row_label = f"category subrow {predictor!r}/{level.get('level', '')!r}"
        expected_level_values = {
            "Stage": stage_map.get(predictor, ""),
            "Effect representation / comparison": level.get(
                "effect_representation", ""
            ),
            "N": level.get("level_n", ""),
            "Events": level.get("level_events", ""),
            "Non-events": level.get("level_non_events", ""),
            "Unique subjects": level.get("level_unique_subjects", ""),
            "Crude OR": level.get("odds_ratio", ""),
            "95% CI": _format_ci(level, "ci_low", "ci_high"),
            "Level p": level.get("coef_p_value", ""),
            "Fit status": level.get("fit_status", ""),
            "Separation status": level.get("separation_status", ""),
            "Support status": level.get("support_status", ""),
            "Reason": level.get("reason", ""),
        }
        for column, expected in expected_level_values.items():
            _assert_value_equal(
                display[column], expected, row_label=row_label, column=column
            )


def save_appendix_table_a2_xlsx(
    appendix_table_a2: pd.DataFrame,
    filename: str = exports.APPENDIX_TABLE_A2_FULL_UNIVARIABLE_XLSX,
) -> Path:
    """Write a lightly formatted XLSX suitable for book-facing review."""
    path = exports.assert_safe_output_path(exports.TABLES_DIR / filename)
    path.parent.mkdir(parents=True, exist_ok=True)
    appendix_table_a2.to_excel(path, index=False)

    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    workbook = load_workbook(path)
    worksheet = workbook.active
    worksheet.freeze_panes = "A2"
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    source_fill = PatternFill("solid", fgColor="F2F2F2")
    subrow_fill = PatternFill("solid", fgColor="FFFFFF")

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(wrap_text=True)

    columns = {cell.value: cell.column for cell in worksheet[1]}
    for row_idx in range(2, worksheet.max_row + 1):
        row_type = worksheet.cell(row_idx, columns["Row type"]).value
        fill = source_fill if row_type == "source" else subrow_fill
        for col_idx in range(1, worksheet.max_column + 1):
            worksheet.cell(row_idx, col_idx).fill = fill
        if row_type == "category_subrow":
            worksheet.cell(row_idx, columns["Level"]).alignment = Alignment(indent=1)
            worksheet.cell(row_idx, columns["Clinical Label"]).font = Font(italic=True)

    width_by_name = {
        "Variable": 32,
        "Clinical Label": 42,
        "Row type": 16,
        "Level": 38,
        "Reference": 32,
        "Stage": 10,
        "Effect representation / comparison": 34,
        "N": 12,
        "Events": 12,
        "Non-events": 12,
        "Unique subjects": 15,
        "Crude OR": 14,
        "95% CI": 28,
        "Predictor-level raw p": 18,
        "Predictor-level BH q": 18,
        "Level p": 14,
        "In BH family": 14,
        "Fit status": 28,
        "Separation status": 22,
        "Support status": 30,
        "Reason": 80,
    }
    for header, col_idx in columns.items():
        worksheet.column_dimensions[get_column_letter(col_idx)].width = width_by_name.get(
            header, 18
        )
    for header in ("Predictor-level raw p", "Predictor-level BH q", "Level p"):
        col_idx = columns[header]
        for row_idx in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row_idx, col_idx)
            if cell.value not in (None, ""):
                try:
                    cell.value = float(cell.value)
                except (TypeError, ValueError):
                    pass
            cell.number_format = APPENDIX_TABLE_A2_P_VALUE_NUMBER_FORMAT
    workbook.save(path)
    return path


def save_appendix_table_a2_csv(
    appendix_table_a2: pd.DataFrame,
    filename: str = exports.APPENDIX_TABLE_A2_FULL_UNIVARIABLE_CSV,
) -> Path:
    path = exports.assert_safe_output_path(exports.TABLES_DIR / filename)
    path.parent.mkdir(parents=True, exist_ok=True)
    appendix_table_a2.to_csv(path, index=False, float_format="%.17g")
    return path


def _appendix_a2_markdown_text(appendix_table_a2: pd.DataFrame) -> str:
    def escape(value: object) -> str:
        if pd.isna(value):
            return ""
        text = str(value)
        text = text.replace("\\", "\\\\")
        text = text.replace("|", "\\|")
        text = text.replace("\r\n", "<br>").replace("\n", "<br>").replace("\r", "<br>")
        return text

    headers = [str(c) for c in appendix_table_a2.columns]
    lines = [
        "| " + " | ".join(escape(h) for h in headers) + " |",
        "| " + " | ".join("---" for _ in headers) + " |",
    ]
    for _, row in appendix_table_a2.iterrows():
        lines.append("| " + " | ".join(escape(v) for v in row) + " |")
    return "\n".join(lines) + "\n"


def save_appendix_table_a2_markdown(
    appendix_table_a2: pd.DataFrame,
    filename: str = exports.APPENDIX_TABLE_A2_FULL_UNIVARIABLE_MD,
) -> Path:
    path = exports.assert_safe_output_path(exports.TABLES_DIR / filename)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(_appendix_a2_markdown_text(appendix_table_a2), encoding="utf-8")
    return path


def save_appendix_table_a2_exports(
    appendix_table_a2: pd.DataFrame,
    *,
    include_markdown: bool = True,
    csv_filename: str = exports.APPENDIX_TABLE_A2_FULL_UNIVARIABLE_CSV,
    xlsx_filename: str = exports.APPENDIX_TABLE_A2_FULL_UNIVARIABLE_XLSX,
    markdown_filename: str = exports.APPENDIX_TABLE_A2_FULL_UNIVARIABLE_MD,
) -> dict[str, Path]:
    """Write Appendix Table A2 CSV/XLSX and optional Markdown exports."""
    paths = {
        "csv": save_appendix_table_a2_csv(appendix_table_a2, csv_filename),
        "xlsx": save_appendix_table_a2_xlsx(appendix_table_a2, xlsx_filename),
    }
    if include_markdown:
        paths["md"] = save_appendix_table_a2_markdown(
            appendix_table_a2, markdown_filename
        )
    return paths
